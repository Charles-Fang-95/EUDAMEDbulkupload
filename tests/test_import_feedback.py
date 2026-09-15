import tempfile
import unittest
from pathlib import Path
from datetime import datetime
from local_beta.importer import WorkbookImporter
from local_beta import views
import openpyxl


class ImportFeedbackTests(unittest.TestCase):
    def test_sparse_loading_preserves_values_types_and_number_formats(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'sample.xlsx'
            wb = openpyxl.Workbook()
            ws = wb.active
            for row, value in [(1, 'Header'), (4, '00123'), (12, 123), (20, datetime(2026, 9, 15)), (40, '=literal')]:
                ws.cell(row, 1, value)
            ws['A40'].data_type = 's'
            ws['A12'].number_format = '000000'
            ws['A10000'].number_format = '000000'
            wb.save(path)
            expected = openpyxl.load_workbook(path, data_only=True)
            actual = WorkbookImporter._load_import_workbook(path)
            for row in [1, 4, 12, 20, 40]:
                a, b = actual.active.cell(row,1), expected.active.cell(row,1)
                self.assertEqual((a.value,a.data_type,a.number_format), (b.value,b.data_type,b.number_format))
            self.assertEqual(actual.active.max_row, 40)

    def test_import_shows_busy_message_and_result_before_form(self):
        html = views.import_page()
        self.assertIn('id="import-progress"', html)
        self.assertIn("disabled=true", html)
        result = dict(validation={'errors':[], 'warnings':[]}, summary={'basic_count':1,'udi_count':1}, import_id=1)
        html = views.import_page('完成', result, 'success')
        self.assertLess(html.index('id="import-result"'), html.index('action="/import"'))
        self.assertIn('result.scrollIntoView', html)

    def test_precheck_errors_and_warnings_are_distinct_and_persistent(self):
        html = views.export_result_panel(dict(service_type='DEVICE.POST',action='preflight',errors=['Missing code'],warnings=['Check country']), 'DEVICE.POST')
        self.assertIn('class="alert error"',html)
        self.assertIn('class="alert warning"',html)
        self.assertIn('Missing code',html)
        self.assertNotIn('window.location.assign',html)
