"""Synthetic BR-UDID-131 regression cases; no customer identifiers or XML."""
import unittest
from pathlib import Path
from unittest.mock import patch

from local_beta.importer import WorkbookImporter, DataValidator
from local_beta import importer as importer_module
from local_beta.template_schema import columns_for_entry_sheet, TEMPLATE_VERSION
from openpyxl import Workbook
from test_xsd_validation import XSDValidationBase, ET, _schema

SERVICES = ("DEVICE.POST", "UDI_DI.POST", "UDI_DI.PATCH")


def write_kit_workbook(path, description="", language="", kit=True, profile="IVDD"):
    wb = Workbook()
    ws = wb.active
    ws.title = profile
    values = {
        "basic": {"Basic UDI-DI Code": "SYNTH-BASIC-001", "Issuing Entity": "GS1",
                  "Manufacturer SRN": "DE-MF-000000001", "Risk Class": "IVD General" if profile == "IVDD" else "Class B",
                  "Applicable Legislation": profile, "Device Type": "Regular Device",
                  "Device Name/Model": "Synthetic kit", "EMDN Code": "W0101", "Is it a Kit": kit},
        "udi": {"UDI-DI Code": "SYNTH-UDI-001", "UDI-DI Issuing Entity": "GS1",
                "Device Status": "No longer placed on the EU market", "Reference Number": "SYNTH-REF-001",
                "Single Use Device": False, "Device Labelled as Sterile": False, "Trade Name Applicable": False,
                "Nomenclature Code": "W0101", "Additional Description": description, "Description Language": language},
    }
    for i, col in enumerate(columns_for_entry_sheet(profile), 1):
        ws.cell(1, i, col["header"])
        value = values.get(col["entity"], {}).get(col["field"])
        if value is not None:
            ws.cell(4, i, value)
    wb.create_sheet("How to Use").cell(1, 1, f"EUDAMED Template {TEMPLATE_VERSION} - How to Use")
    wb.save(path)


class DescriptionRuleTests(XSDValidationBase):
    def test_missing_description_blocks_all_three_services(self):
        for profile in ("IVDD", "IVDR", "PR"):
            for kit in (True, "TRUE", " true "):
                with self.subTest(profile=profile, kit=kit):
                    _, uid = self._seed(profile, version="1", basic_over={"Is it a Kit": kit},
                                        udi_over={"Additional Description": " \t"})
                    item = self.repo.get_udis_by_ids([uid])[0]
                    validator = DataValidator({"Basic UDI-DI": [item["basic_payload"]], "UDI-DI": [item["payload"]]})
                    validator._layer2_completeness_validation()
                    self.assertTrue(any(e.error_type == "CONDITIONAL_REQUIRED_MISSING" and e.field == "Additional Description" for e in validator.errors))
                    for service in SERVICES:
                        result = self.exporter.export(service, [uid])
                        self.assertTrue(any("BR-UDID-131" in e for e in result["errors"]))
                        self.assertNotIn("file_path", result)
        self.assertEqual(list((self.tmp / "exports").iterdir()), [])

    def test_system_and_procedure_pack_each_require_description(self):
        for device_type in ("System", "Procedure Pack"):
            _, uid = self._seed("PR", version="1", basic_over={"Device Type": device_type},
                                udi_over={"Additional Description": ""})
            for service in SERVICES:
                self.assertTrue(any("BR-UDID-131" in e for e in self.exporter.validate(service, [uid])["errors"]))

    def test_any_description_requires_explicit_language(self):
        for profile in ("IVDD", "IVDR", "MDR", "PR"):
            for kit in (True, False):
                _, uid = self._seed(profile, version="1", basic_over={"Is it a Kit": kit},
                                    udi_over={"Additional Description": "Synthetic description", "Description Language": " "})
                for service in SERVICES:
                    result = self.exporter.export(service, [uid])
                    self.assertTrue(any("Description Language" in e for e in result["errors"]))
                    self.assertNotIn("file_path", result)

    def test_valid_description_preserved_and_xsd_passes(self):
        for profile in ("IVDD", "IVDR", "PR"):
            _, uid = self._seed(profile, version="1", basic_over={"Is it a Kit": True},
                                udi_over={"Additional Description": "Synthetic kit description", "Description Language": "de"})
            for service in SERVICES:
                for xml in self._xmls(service, [uid]):
                    doc = ET.fromstring(xml)
                    self.assertTrue(_schema().validate(doc), str(_schema().error_log))
                    descriptions = doc.xpath('//*[local-name()="additionalDescription"]')
                    self.assertEqual(len(descriptions), 1)
                    self.assertEqual(descriptions[0].xpath('.//*[local-name()="language"]/text()'), ["DE"])
                    self.assertEqual(descriptions[0].xpath('.//*[local-name()="textValue"]/text()'), ["Synthetic kit description"])

    def test_ordinary_regulation_and_legacy_remain_optional(self):
        for profile in ("MDR", "IVDR", "MDD", "AIMDD", "IVDD"):
            for kit in (False, "FALSE", " false "):
                _, uid = self._seed(profile, version="1", basic_over={"Is it a Kit": kit})
                for service in SERVICES:
                    self.assertEqual(self.exporter.validate(service, [uid])["errors"], [])
                item = self.repo.get_udis_by_ids([uid])[0]
                validator = DataValidator({"Basic UDI-DI": [item["basic_payload"]], "UDI-DI": [item["payload"]]})
                validator._layer2_completeness_validation()
                self.assertFalse(any(e.field in {"Additional Description", "Description Language"} for e in validator.errors))

    def test_excel_false_is_not_reported_as_blank(self):
        self.assertFalse(self.exporter._is_blank_bool(False))
        self.assertFalse(self.exporter._is_blank_bool("FALSE"))
        self.assertTrue(self.exporter._is_blank_bool(None))
        self.assertTrue(self.exporter._is_blank_bool(" "))

    def test_patch_version_is_still_required(self):
        _, uid = self._seed("IVDD", basic_over={"Is it a Kit": True},
                            udi_over={"Additional Description": "Synthetic kit", "Description Language": "en"})
        self.assertTrue(any("版本号" in e for e in self.exporter.validate("UDI_DI.PATCH", [uid])["errors"]))

    def test_explicit_false_is_not_overridden_by_old_alias(self):
        _, uid = self._seed("IVDR", basic_over={"Is it a Kit": False, "Kit (IVDR)": True})
        for xml in self._xmls("DEVICE.POST", [uid]):
            doc = ET.fromstring(xml)
            self.assertEqual(doc.xpath('//*[local-name()="kit"]/text()'), ["false"])
            self.assertTrue(_schema().validate(doc), str(_schema().error_log))
        _, alias_uid = self._seed("IVDR", basic_over={"Is it a Kit": "", "Kit (IVDR)": True})
        self.assertTrue(any("BR-UDID-131" in e for e in self.exporter.validate("DEVICE.POST", [alias_uid])["errors"]))

    def test_excel_import_saves_records_with_conditional_errors_then_corrects(self):
        with patch.object(importer_module, "EXPORT_DIR", self.tmp / "exports"):
            for profile in ("IVDD", "IVDR"):
                for kit in (True, "TRUE"):
                    path = self.tmp / "kit.xlsx"
                    write_kit_workbook(path, kit=kit, profile=profile)
                    result = WorkbookImporter(self.repo).import_workbook(path)
                    self.assertEqual(result["summary"]["udi_count"], 1)
                    self.assertTrue(any(e["error_type"] == "CONDITIONAL_REQUIRED_MISSING" and e["field"] == "Additional Description"
                                        for e in result["validation"]["errors"]))
                    uid = next(u["id"] for u in self.repo.list_udis(limit=999) if u["udi_code"] == "SYNTH-UDI-001")
                    self.assertTrue(self.exporter.validate("DEVICE.POST", [uid])["errors"])
                    write_kit_workbook(path, description="Synthetic diagnostic kit", profile=profile)
                    result = WorkbookImporter(self.repo).import_workbook(path)
                    self.assertTrue(any(e["error_type"] == "CONDITIONAL_REQUIRED_MISSING" and e["field"] == "Description Language"
                                        for e in result["validation"]["errors"]))
                    write_kit_workbook(path, description="Synthetic diagnostic kit", language="en", profile=profile)
                    result = WorkbookImporter(self.repo).import_workbook(path)
                    self.assertEqual(result["validation"]["errors"], [])
                    self._assert_valid("DEVICE.POST", [uid], profile)
