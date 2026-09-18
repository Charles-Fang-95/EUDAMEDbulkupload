import unittest
from pathlib import Path
from local_beta import template_schema as schema
from local_beta.importer import WorkbookImporter
from local_beta.template_migrator import _normalize_main_row
from tests.test_xsd_validation import XSDValidationBase, _schema
from lxml import etree as ET
import openpyxl

ALIASES = ['Orthopedic', 'Orthopaedic', 'ORTHOPEDIC', 'MDR_ORTHOPEDIC', 'MDD_ORTHOPEDIC', 'AIMDD_ORTHOPEDIC', 'MDR_ORTHOPEDIC - Orthopedic']

class RetiredOrthopedicTests(XSDValidationBase):
    def test_old_database_values_block_registration_and_basic_update(self):
        for value in ALIASES:
            with self.subTest(value=value):
                bid, uid = self._seed('MDR', version='1', basic_over={'Special Device Type':value})
                for service, ids in [('DEVICE.POST',[uid]),('Basic_UDI.PATCH',[bid])]:
                    result = self.exporter.export(service,ids)
                    self.assertTrue(any('Orthopedic' in e and '清空' in e for e in result['errors']),result)
                    self.assertFalse(result.get('file_path'))
                self.assertEqual(self.repo.get_basic_by_code(self.repo.get_udis_by_ids([uid])[0]['basic_code'])['payload']['Special Device Type'], value)

    def test_blank_and_software_still_export_and_validate(self):
        for value in ['', 'Software', 'MDR_SOFTWARE']:
            _, uid=self._seed('MDR',basic_over={'Special Device Type':value})
            result=self.exporter.export('DEVICE.POST',[uid])
            self.assertFalse(result['errors'], result['errors'])
            doc=ET.parse(result['file_path']);_schema().assertValid(doc)
            nodes=doc.xpath('//*[local-name()="specialDevice"]/text()')
            self.assertEqual(nodes, ['MDR_SOFTWARE'] if value else [])

    def test_import_validation_and_migration_preserve_but_explain_old_value(self):
        for value in ALIASES:
            errors=[]
            WorkbookImporter(None)._validate_main_field_rules({'Basic UDI-DI':[{'Basic UDI-DI Code':'B1','Applicable Legislation':'MDR','Special Device Type':value}]},errors)
            self.assertTrue(any('Orthopedic' in e['message'] for e in errors))
            wb=openpyxl.Workbook();ws=wb.active;ws.title='MDR'
            headers=[c['header'] for c in schema.columns_for_entry_sheet('MDR')]
            col=headers.index('Basic - Special Device Type')+1
            ws.cell(4,col,value);report={'warnings':[]}
            _normalize_main_row(ws,headers,4,report)
            self.assertEqual(ws.cell(4,col).value,value)
            self.assertTrue(any('Orthopedic' in w for w in report['warnings']))

    def test_template_choices_exclude_only_retired_codes(self):
        raw=schema._xsd_enum_values_from_file(schema.BASIC_UDI_XSD,'MDRSpecialDeviceTypeEnum')
        current=schema.ENUM_SOURCES['special_device_mdr']
        self.assertEqual([v for v in raw if v not in current], [v for v in raw if schema.is_retired_orthopedic(v)])
        self.assertEqual(len(raw)-len(current),3)
        for key in ['special_device_mdr','special_device_ivdr']:
            self.assertFalse(any(schema.is_retired_orthopedic(v) for v in schema.ENUM_SOURCES[key]))
